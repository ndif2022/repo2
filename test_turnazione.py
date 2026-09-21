# -*- coding: utf-8 -*-
"""
Verifica del file Excel della turnazione.

Compila il foglio «Assenze» con scenari noti, ricalcola il file con LibreOffice
Calc e controlla che turni, sostituzioni, doppi turni e recuperi rispettino le
regole R1-R7 descritte nel foglio «Regole».

Requisiti: openpyxl, libreoffice-calc.
Uso: python3 test_turnazione.py [file.xlsx]
"""

import datetime as dt
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

LICENZA = "Licenza - settimana intera"
ESTERNA = "Attività esterna - settimana intera"
ALTRO = "Altro motivo"

TUR_R0 = 6          # prima riga dati del foglio Turnazione
ASS_R0 = 5          # prima riga dati del foglio Assenze
RIEP_R0 = 6         # prima riga dati del foglio Riepilogo

esiti = {"ok": 0, "ko": 0}


def verifica(descrizione, atteso, ottenuto):
    if atteso == ottenuto:
        esiti["ok"] += 1
        print(f"  [OK] {descrizione} -> {ottenuto!r}")
    else:
        esiti["ko"] += 1
        print(f"  [KO] {descrizione}\n       atteso   = {atteso!r}\n       ottenuto = {ottenuto!r}")


class Banco:
    """Prepara una copia del file, vi scrive le assenze e la fa ricalcolare."""

    def __init__(self, sorgente, lavoro):
        self.sorgente = sorgente
        self.lavoro = lavoro

    def esegui(self, tag, assenze, settimana_comunicazione=None):
        copia = os.path.join(self.lavoro, f"{tag}.xlsx")
        wb = load_workbook(self.sorgente)
        ws = wb["Assenze"]
        for i, (sett, persona, tipo, data) in enumerate(assenze):
            r = ASS_R0 + i
            ws.cell(r, 1).value = sett
            ws.cell(r, 3).value = persona
            ws.cell(r, 4).value = tipo
            ws.cell(r, 5).value = data
        if settimana_comunicazione is not None:
            wb["Comunicazione Giovedì"]["C4"] = settimana_comunicazione
        wb.save(copia)

        uscita = os.path.join(self.lavoro, f"out_{tag}")
        subprocess.run(
            ["soffice", "--headless", "--norestore",
             f"-env:UserInstallation=file://{self.lavoro}/profilo_{tag}",
             "--convert-to", "xlsx", "--outdir", uscita, copia],
            check=True, capture_output=True, timeout=300,
        )
        return load_workbook(os.path.join(uscita, f"{tag}.xlsx"), data_only=True)


def calendario(wb, settimane=13):
    ws = wb["Turnazione"]
    righe = []
    for i in range(settimane * 5):
        r = TUR_R0 + i
        righe.append({
            "sett": ws.cell(r, 1).value, "data": ws.cell(r, 2).value,
            "giorno": ws.cell(r, 3).value, "previsto": ws.cell(r, 4).value,
            "stato": ws.cell(r, 5).value, "conducente": ws.cell(r, 6).value,
            "tipo": ws.cell(r, 7).value, "riposo": ws.cell(r, 8).value,
            "motivo": ws.cell(r, 9).value,
        })
    return righe


def riepilogo(wb):
    ws = wb["Riepilogo"]
    return {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(2, 9)]
            for r in range(RIEP_R0, RIEP_R0 + 5)}


def main():
    sorgente = sys.argv[1] if len(sys.argv) > 1 else \
        os.path.join(os.path.dirname(os.path.abspath(__file__)),
                     "turnazione-auto-trinitapoli-foggia-20260921.xlsx")
    if not os.path.exists(sorgente):
        raise SystemExit(f"File non trovato: {sorgente}")
    if not shutil.which("soffice"):
        raise SystemExit("LibreOffice Calc non disponibile: impossibile ricalcolare il file.")

    lavoro = tempfile.mkdtemp(prefix="turnazione_test_")
    os.environ.setdefault("HOME", lavoro)
    banco = Banco(sorgente, lavoro)

    # ------------------------------------------------------------- SCENARIO A
    print("\nSCENARIO A - nessuna assenza: turnazione base")
    cal = calendario(banco.esegui("a", []))
    verifica("65 giorni di calendario", 65, len(cal))
    verifica("tutti i turni sono ordinari", {"Ordinario"}, {x["tipo"] for x in cal})
    verifica("primo turno: lunedì 28/09/2026 Nicola",
             ("Nicola", dt.datetime(2026, 9, 28)), (cal[0]["conducente"], cal[0]["data"]))
    verifica("ultimo turno: venerdì 25/12/2026 Savino",
             ("Savino", dt.datetime(2026, 12, 25)), (cal[64]["conducente"], cal[64]["data"]))
    conta = {}
    for x in cal:
        conta[x["conducente"]] = conta.get(x["conducente"], 0) + 1
    verifica("13 turni a testa",
             {"Nicola": 13, "Stefano": 13, "Rocco": 13, "Giuseppe": 13, "Savino": 13}, conta)

    # ------------------------------------------------------------- SCENARIO B
    print("\nSCENARIO B - R1: licenza di settimana intera, turno NON da recuperare")
    wb = banco.esegui("b", [(2, "Nicola", LICENZA, None)])
    cal = calendario(wb)
    lunedi2 = cal[5]
    verifica("stato del titolare", LICENZA, lunedi2["stato"])
    verifica("tipo di turno", "Sostituzione (doppio turno)", lunedi2["tipo"])
    verifica("sostituto = Stefano (successivo a Nicola)", "Stefano", lunedi2["conducente"])
    verifica("il motivo dichiara il turno non recuperabile",
             True, "non da recuperare" in (lunedi2["motivo"] or ""))
    verifica("la settimana 3 resta ordinaria", {"Ordinario"}, {x["tipo"] for x in cal[10:15]})
    r = riepilogo(wb)
    verifica("Nicola: 12 ordinari, 0 doppi, 0 recuperi, 12 guide, 1 saltato non recuperabile",
             [12, 0, 0, 12, 1, 0, 0], r["Nicola"])
    verifica("Stefano: 13 ordinari, 1 doppio turno, 14 guide", [13, 1, 0, 14, 0, 0, 0], r["Stefano"])

    # ------------------------------------------------------------- SCENARIO C
    print("\nSCENARIO C - R3+R5: assenza per altro motivo e recupero la settimana dopo")
    wb = banco.esegui("c", [(2, "Nicola", LICENZA, None),
                            (3, "Rocco", ALTRO, dt.date(2026, 10, 14))])
    cal = calendario(wb)
    mercoledi3, giovedi4 = cal[12], cal[18]
    verifica("mer 14/10 sett.3: titolare Rocco", ("Rocco", dt.datetime(2026, 10, 14)),
             (mercoledi3["previsto"], mercoledi3["data"]))
    verifica("mer sett.3: sostituto = Giuseppe (Stefano ha già un doppio turno)",
             ("Sostituzione (doppio turno)", "Giuseppe"),
             (mercoledi3["tipo"], mercoledi3["conducente"]))
    verifica("mer sett.3: recupero annunciato a partire dalla settimana 4",
             True, "a partire dalla settimana 4" in (mercoledi3["motivo"] or ""))
    verifica("gio sett.4: Rocco recupera sul turno di Giuseppe",
             ("Recupero", "Rocco", "Giuseppe"),
             (giovedi4["tipo"], giovedi4["conducente"], giovedi4["riposo"]))
    verifica("la settimana 5 torna ordinaria", {"Ordinario"}, {x["tipo"] for x in cal[20:25]})
    r = riepilogo(wb)
    verifica("Rocco: 1 turno da recuperare, 1 recupero effettuato, 13 guide",
             (1, 1, 13), (r["Rocco"][5], r["Rocco"][2], r["Rocco"][3]))
    verifica("Giuseppe: bilancio in pari dopo il doppio turno", [12, 1, 0, 13, 0, 0, 0],
             r["Giuseppe"])

    # ------------------------------------------------------------- SCENARIO D
    print("\nSCENARIO D - R4: i doppi turni ruotano su persone diverse")
    cal = calendario(banco.esegui("d", [(2, "Nicola", LICENZA, None),
                                        (3, "Nicola", ESTERNA, None),
                                        (4, "Nicola", LICENZA, None)]))
    sostituti = [cal[5]["conducente"], cal[10]["conducente"], cal[15]["conducente"]]
    verifica("tre sostituti diversi", 3, len(set(sostituti)))
    verifica("ordine dei sostituti", ["Stefano", "Rocco", "Giuseppe"], sostituti)
    verifica("nessun recupero generato", {"Ordinario", "Sostituzione (doppio turno)"},
             {x["tipo"] for x in cal})

    # ------------------------------------------------------------- SCENARIO E
    print("\nSCENARIO E - R7: il recupero slitta di una settimana")
    # Rocco salta mer sett.2; Giuseppe lo sostituisce e diventa creditore.
    # In sett.3 Rocco e' in licenza: il recupero non puo' avvenire e slitta alla sett.4.
    wb = banco.esegui("e", [(2, "Rocco", ALTRO, dt.date(2026, 10, 7)),
                            (3, "Rocco", LICENZA, None)])
    cal = calendario(wb)
    verifica("mer sett.2: sostituto = Giuseppe", "Giuseppe", cal[7]["conducente"])
    verifica("gio sett.3: il recupero non scatta, Rocco è assente",
             ("Ordinario", "Giuseppe"), (cal[13]["tipo"], cal[13]["conducente"]))
    verifica("gio sett.4: il recupero slitta qui",
             ("Recupero", "Rocco", "Giuseppe"),
             (cal[18]["tipo"], cal[18]["conducente"], cal[18]["riposo"]))
    verifica("gio sett.4: il motivo cita la settimana di origine del debito",
             True, "saltato nella settimana 2" in (cal[18]["motivo"] or ""))
    verifica("la settimana 5 torna ordinaria", {"Ordinario"}, {x["tipo"] for x in cal[20:25]})
    r = riepilogo(wb)
    verifica("Rocco: 1 da recuperare, 1 recupero effettuato, 0 ancora da saldare",
             (1, 1, 0), (r["Rocco"][5], r["Rocco"][2], r["Rocco"][6]))
    verifica("Rocco: 1 turno saltato non recuperabile (licenza sett.3)", 1, r["Rocco"][4])
    verifica("Giuseppe: bilancio in pari (1 doppio turno, 1 riposo compensativo)",
             13, r["Giuseppe"][3])

    # ------------------------------------------------------------- SCENARIO F
    print("\nSCENARIO F - R7: il recupero slitta di più settimane")
    # Rocco salta mer sett.2 e resta assente per tutte le sett.3 e 4:
    # il debito verso Giuseppe si salda alla prima settimana utile, la 5.
    wb = banco.esegui("f", [(2, "Rocco", ALTRO, dt.date(2026, 10, 7)),
                            (3, "Rocco", LICENZA, None),
                            (4, "Rocco", ESTERNA, None)])
    cal = calendario(wb)
    verifica("gio sett.3: nessun recupero", "Ordinario", cal[13]["tipo"])
    verifica("gio sett.4: nessun recupero", "Ordinario", cal[18]["tipo"])
    verifica("gio sett.5: il recupero si esegue qui",
             ("Recupero", "Rocco", "Giuseppe"),
             (cal[23]["tipo"], cal[23]["conducente"], cal[23]["riposo"]))
    verifica("le settimane 6 e 7 sono ordinarie", {"Ordinario"}, {x["tipo"] for x in cal[25:35]})
    r = riepilogo(wb)
    verifica("Rocco: 1 recupero effettuato, 0 ancora da saldare",
             (1, 0), (r["Rocco"][2], r["Rocco"][6]))
    verifica("Rocco: 2 turni saltati non recuperabili (sett.3 e 4)", 2, r["Rocco"][4])

    # ------------------------------------------------------------- SCENARIO G
    print("\nSCENARIO G - catena di due assenze recuperabili consecutive")
    wb = banco.esegui("g", [(2, "Nicola", ALTRO, dt.date(2026, 10, 5)),
                            (3, "Nicola", ALTRO, dt.date(2026, 10, 12))])
    cal = calendario(wb)
    verifica("lun sett.2: sostituto = Stefano", "Stefano", cal[5]["conducente"])
    verifica("mar sett.3: Nicola restituisce a Stefano",
             ("Recupero", "Nicola", "Stefano"),
             (cal[11]["tipo"], cal[11]["conducente"], cal[11]["riposo"]))
    verifica("lun sett.3: Nicola assente di nuovo, sostituto diverso da Stefano",
             ("Sostituzione (doppio turno)", "Rocco"), (cal[10]["tipo"], cal[10]["conducente"]))
    verifica("mer sett.4: Nicola restituisce a Rocco",
             ("Recupero", "Nicola", "Rocco"),
             (cal[17]["tipo"], cal[17]["conducente"], cal[17]["riposo"]))
    r = riepilogo(wb)
    verifica("Nicola: 2 da recuperare, 2 recuperi effettuati, 0 da saldare",
             (2, 2, 0), (r["Nicola"][5], r["Nicola"][2], r["Nicola"][6]))
    verifica("tutti chiudono con 13 guide",
             [13, 13, 13, 13, 13], [r[n][3] for n in
                                    ["Nicola", "Stefano", "Rocco", "Giuseppe", "Savino"]])

    # ------------------------------------------------------------- SCENARIO H
    print("\nSCENARIO H - debito aperto alla fine del periodo")
    wb = banco.esegui("h", [(13, "Savino", ALTRO, dt.date(2026, 12, 25))])
    r = riepilogo(wb)
    verifica("Savino: 1 da recuperare, 0 effettuati, 1 ancora da saldare",
             (1, 0, 1), (r["Savino"][5], r["Savino"][2], r["Savino"][6]))

    # ------------------------------------------------------------- SCENARIO I
    print("\nSCENARIO I - due assenze di settimana intera nella stessa settimana")
    cal = calendario(banco.esegui("i", [(5, "Nicola", LICENZA, None),
                                        (5, "Savino", ESTERNA, None)]))
    sett5 = cal[20:25]
    verifica("tutti e 5 i giorni hanno un conducente", 5,
             len([x for x in sett5 if x["conducente"]]))
    verifica("lunedì e venerdì coperti da sostituti",
             {"Sostituzione (doppio turno)"}, {sett5[0]["tipo"], sett5[4]["tipo"]})
    verifica("i due sostituti sono persone diverse", 2,
             len({sett5[0]["conducente"], sett5[4]["conducente"]}))
    verifica("nessun sostituto è fra gli assenti", set(),
             {sett5[0]["conducente"], sett5[4]["conducente"]} & {"Nicola", "Savino"})
    verifica("nessun recupero nella settimana 6", {"Ordinario"}, {x["tipo"] for x in cal[25:30]})

    # ------------------------------------------------------------- SCENARIO L
    print("\nSCENARIO L - R6: foglio «Comunicazione Giovedì»")
    wb = banco.esegui("l", [(5, "Nicola", LICENZA, None), (5, "Savino", ESTERNA, None)],
                      settimana_comunicazione=5)
    cg = wb["Comunicazione Giovedì"]
    verifica("periodo della settimana 5", "26/10/2026 - 30/10/2026", cg["C5"].value)
    verifica("comunicazione entro giovedì 22/10/2026", True, "22/10/2026" in str(cg["C6"].value))
    verifica("lunedì: conducente e nota sintetica",
             ("Stefano", "in sostituzione di Nicola, doppio turno"),
             (cg.cell(9, 4).value, cg.cell(9, 5).value))
    testo = cg.cell(16, 2).value
    verifica("testo pronto da inviare completo", True,
             isinstance(testo, str) and all(g in testo for g in
                                            ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì"]))
    print("\n--- testo generato per la settimana 5 ---")
    print(testo)

    print(f"\nControlli superati: {esiti['ok']} · falliti: {esiti['ko']}")
    shutil.rmtree(lavoro, ignore_errors=True)
    return 0 if esiti["ko"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
